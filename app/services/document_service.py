import io

import pandas as pd
from docx import Document
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook


class DocumentService:
    
    # --- 1. EXCEL (.xlsx) HANDLING ---
    async def parse_excel(self, file: UploadFile) -> list[dict]:
        """Upload လုပ်လာသော Excel ဖိုင်ကို ဖတ်ပြီး Python List/Dict အဖြစ်သို့ ပြောင်းမည်"""
        try:
            contents = await file.read()
            df = pd.read_excel(io.BytesIO(contents))
            return df.to_dict(orient="records")
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Invalid Excel file: {str(e)}")

    def generate_excel_report(self, data: list[dict], sheet_name: str = "Report") -> io.BytesIO:
        """Data များကို Excel ဖိုင်အဖြစ်သို့ ပြန်လည်ထုတ်ပေးမည် (Download အတွက်)"""
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        if not data:
            raise HTTPException(status_code=400, detail="No data provided for export")

        # Header များကို အလိုအလျောက် ထည့်သွင်းခြင်း
        headers = list(data[0].keys())
        ws.append(headers)

        # Row အလိုက် Data များကို ထည့်သွင်းခြင်း
        for row in data:
            ws.append([row.get(h, "") for h in headers])

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    # --- 2. CSV HANDLING ---
    async def parse_csv(self, file: UploadFile) -> list[dict]:
        """Upload လုပ်လာသော CSV ဖိုင်ကို ဖတ်မည်"""
        try:
            contents = await file.read()
            df = pd.read_csv(io.BytesIO(contents))
            return df.to_dict(orient="records")
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Invalid CSV file: {str(e)}")

    # --- 3. WORD (.docx) HANDLING ---
    def generate_word_report(self, title: str, content_paragraphs: list[str]) -> io.BytesIO:
        """Word Report တစ်စောင်ကို အလိုအလျောက် တည်ဆောက်ပေးမည်"""
        doc = Document()
        
        # Title ထည့်ခြင်း
        doc.add_heading(title, level=1)
        
        # Paragraph များ ထည့်ခြင်း
        for para in content_paragraphs:
            doc.add_paragraph(para)
            
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        return output

document_service = DocumentService()