import csv
import io
import logging
import math
import os
import re
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from typing import Annotated, Any
from uuid import UUID
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from fastapi import (
    APIRouter,
    Depends,
    File,
    Form,
    HTTPException,
    Query,
    Request,
    UploadFile,
    status,
)
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from slowapi import Limiter
from slowapi.util import get_remote_address
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.dependencies import (
    RoleChecker,
    enforce_permission,
    get_authorized_department_ids,
)
from app.core.timezone import as_application_time
from app.database.connection import get_db
from app.modules.departments.models import UserDepartment
from app.modules.kpi.engine import KPICalculationEngine
from app.modules.kpi.models import (
    KPIFormula,
    KPIPeriod,
    KPIPeriodStatus,
    KPIRecord,
    PeriodType,
)
from app.modules.kpi.period_ranges import calendar_period_range
from app.modules.kpi.upload_contract import (
    EMPLOYEE_HEADERS,
    HeaderContract,
    HeaderContractError,
    MetricHeader,
    build_header_contract,
    normalize_header,
)
from app.modules.users.models import User

logger = logging.getLogger(__name__)
limiter = Limiter(key_func=get_remote_address)
router = APIRouter(prefix="/kpi-uploads", tags=["KPI Excel Uploads"])
admin_hr_checker = RoleChecker(["Super Admin", "Admin", "HR"])

MAX_FILE_SIZE = 5 * 1024 * 1024
ALLOWED_EXTENSIONS = {".csv", ".xlsx"}
EDITABLE_PERIOD_STATUSES = {KPIPeriodStatus.OPEN, KPIPeriodStatus.IN_PROGRESS}


class UploadValidationError(ValueError):
    def __init__(self, errors: list[str] | str):
        self.errors = [errors] if isinstance(errors, str) else errors
        super().__init__("KPI workbook validation failed.")


@dataclass(frozen=True)
class ParsedWorkbook:
    headers: list[Any]
    rows: list[tuple[int, list[Any]]]


@dataclass(frozen=True)
class ValidatedUploadRow:
    row_number: int
    email: str
    user: User
    formula: KPIFormula
    task_data: dict[str, float]


def _period_month_range(period: KPIPeriod):
    start_date = as_application_time(period.start_date).date()
    if start_date.day != 1:
        return None
    expected = calendar_period_range(PeriodType.MONTHLY, start_date)
    if as_application_time(period.start_date) != expected.start_inclusive or (
        as_application_time(period.end_date) != expected.end_exclusive
    ):
        return None
    return expected


def _assert_importable_period(period: KPIPeriod | None) -> KPIPeriod:
    if period is None or period.is_deleted:
        raise UploadValidationError("Select an available evaluation period.")
    if period.status not in EDITABLE_PERIOD_STATUSES:
        raise UploadValidationError(
            "The selected evaluation period is not open for KPI imports."
        )
    if _period_month_range(period) is None:
        raise UploadValidationError(
            "The selected evaluation period must use complete monthly boundaries."
        )
    return period


async def _resolve_period(
    db: AsyncSession,
    *,
    period_id: UUID | None,
    target_month: str | None,
) -> KPIPeriod:
    if period_id is not None:
        period = await db.get(KPIPeriod, period_id)
        return _assert_importable_period(period)

    if target_month is None or not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", target_month):
        raise UploadValidationError(
            "Select an evaluation period using its period ID."
        )
    parsed_date = date.fromisoformat(f"{target_month}-01")
    expected = calendar_period_range(PeriodType.MONTHLY, parsed_date)
    candidates = (
        await db.execute(
            select(KPIPeriod).where(
                KPIPeriod.is_deleted.is_(False),
                KPIPeriod.status.in_(EDITABLE_PERIOD_STATUSES),
            )
        )
    ).scalars().all()
    matches = [
        period
        for period in candidates
        if _period_month_range(period) == expected
    ]
    if len(matches) != 1:
        raise UploadValidationError(
            "The selected month must match exactly one open evaluation period."
        )
    return _assert_importable_period(matches[0])


async def _resolve_user_formula(
    db: AsyncSession, user: User, period: KPIPeriod
) -> KPIFormula | None:
    period_range = _period_month_range(period)
    try:
        formula, resolved_period = (
            await KPICalculationEngine.get_applicable_formula_and_period(
                db, user, period_range.start_inclusive
            )
        )
    except ValueError as exc:
        raise UploadValidationError(
            f"KPI formula resolution is ambiguous for {user.email} in the selected period."
        ) from exc
    if resolved_period is None or resolved_period.id != period.id:
        raise UploadValidationError(
            f"No unambiguous KPI period/formula context exists for {user.email}."
        )
    return formula


async def _read_upload_dataframe(file: UploadFile) -> ParsedWorkbook:
    filename = file.filename or ""
    extension = os.path.splitext(filename)[1].lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise UploadValidationError(
            "Unsupported file format. Upload a .xlsx or .csv file."
        )
    file.file.seek(0, os.SEEK_END)
    file_size = file.file.tell()
    await file.seek(0)
    if file_size > MAX_FILE_SIZE:
        raise UploadValidationError("File exceeds the 5MB upload limit.")
    contents = await file.read()
    if len(contents) > MAX_FILE_SIZE:
        raise UploadValidationError("File exceeds the 5MB upload limit.")
    if not contents:
        raise UploadValidationError("The uploaded workbook is empty.")

    def parse_file() -> ParsedWorkbook:
        if extension == ".csv":
            return _parse_csv(contents)
        return _parse_xlsx(contents)

    try:
        return await run_in_threadpool(parse_file)
    except UploadValidationError:
        raise
    except (UnicodeDecodeError, csv.Error) as exc:
        raise UploadValidationError("The uploaded file is malformed.") from exc


def _parse_csv(contents: bytes) -> ParsedWorkbook:
    try:
        rows = list(csv.reader(io.StringIO(contents.decode("utf-8-sig"), newline="")))
    except (UnicodeDecodeError, csv.Error) as exc:
        raise UploadValidationError("The CSV file is malformed or not UTF-8.") from exc
    if not rows:
        raise UploadValidationError("The uploaded file has no header row.")
    return _table_from_rows(rows)


def _xlsx_row_values(cells, headers, row_number: int) -> list[Any]:
    row_values = [cell.value for cell in cells]
    for column_index, cell in enumerate(cells):
        if cell.data_type == "f":
            header = (
                str(headers[column_index])
                if column_index < len(headers) and headers[column_index]
                else get_column_letter(column_index + 1)
            )
            raise UploadValidationError(
                f"Row {row_number} — {header}: formula cells are not supported."
            )
    return row_values


def _parse_xlsx(contents: bytes) -> ParsedWorkbook:
    try:
        workbook = load_workbook(io.BytesIO(contents), data_only=False, read_only=False)
    except (
        BadZipFile,
        InvalidFileException,
        KeyError,
        OSError,
        ParseError,
        ValueError,
    ) as exc:
        raise UploadValidationError("The XLSX file is corrupted or invalid.") from exc
    try:
        if "KPI_Data" not in workbook.sheetnames:
            raise UploadValidationError(
                "The workbook must contain the KPI_Data worksheet."
            )
        worksheet = workbook["KPI_Data"]
        values = list(worksheet.iter_rows(values_only=False))
        if not values:
            raise UploadValidationError("The KPI_Data worksheet is empty.")
        headers = [cell.value for cell in values[0]]
        rows: list[tuple[int, list[Any]]] = []
        for row_number, cells in enumerate(values[1:], start=2):
            row_values = _xlsx_row_values(cells, headers, row_number)
            if any(value is not None and str(value).strip() for value in row_values):
                rows.append((row_number, row_values))
        return ParsedWorkbook(headers=headers, rows=rows)
    finally:
        workbook.close()


def _table_from_rows(rows: list[list[Any]]) -> ParsedWorkbook:
    headers = rows[0]
    parsed_rows = []
    for row_number, row in enumerate(rows[1:], start=2):
        values = row + [None] * max(0, len(headers) - len(row))
        if any(value is not None and str(value).strip() for value in values):
            parsed_rows.append((row_number, values[: len(headers)]))
    return ParsedWorkbook(headers=headers, rows=parsed_rows)


def _validate_headers(
    workbook_headers: list[Any], contract: HeaderContract
) -> tuple[dict[str, int], list[str]]:
    errors: list[str] = []
    header_indexes: dict[str, int] = {}
    for index, raw_header in enumerate(workbook_headers):
        key = normalize_header(raw_header)
        if not key:
            errors.append(f"Column {index + 1}: header is blank.")
        elif key in header_indexes:
            errors.append(
                f"Duplicate column header after normalization: '{raw_header}'."
            )
        else:
            header_indexes[key] = index

    expected_keys = {normalize_header(header) for header in contract.headers}
    missing_keys = expected_keys - header_indexes.keys()
    unknown_keys = header_indexes.keys() - expected_keys
    canonical_labels = {
        normalize_header(header): header for header in contract.headers
    }
    errors.extend(
        f"Missing required column: {canonical_labels[key]}."
        for key in sorted(missing_keys)
    )
    errors.extend(
        f"Unknown column: {workbook_headers[header_indexes[key]]}."
        for key in sorted(unknown_keys)
    )
    if errors:
        return header_indexes, errors
    return header_indexes, []


def _is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _finite_number(value: Any) -> float | None:
    if isinstance(value, bool) or _is_empty(value):
        return None
    try:
        number = float(Decimal(str(value).strip()))
    except (InvalidOperation, TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _read_upload_metrics(
    row_number: int,
    email: str,
    values: list[Any],
    header_indexes: dict[str, int],
    metric_by_key: dict[str, MetricHeader],
    applicable_keys: frozenset[str],
) -> tuple[dict[str, float], list[str]]:
    task_data: dict[str, float] = {}
    errors: list[str] = []
    for key, field in metric_by_key.items():
        column_index = header_indexes[key]
        value = values[column_index] if column_index < len(values) else None
        if key not in applicable_keys:
            if not _is_empty(value):
                errors.append(
                    f"Row {row_number} — Email '{email}', {field.label}: "
                    "this KPI field does not apply to the employee's formula."
                )
            continue
        number = _finite_number(value)
        if number is None:
            safe_value = repr(value)[:60]
            errors.append(
                f"Row {row_number} — Email '{email}', {field.label}: "
                f"must be a finite number (received {safe_value})."
            )
            continue
        task_data[key] = number
    return task_data, errors


async def _build_upload_rows(
    workbook: ParsedWorkbook | Any,
    db: AsyncSession,
    current_user: User,
    period: KPIPeriod | None = None,
) -> tuple[list[ValidatedUploadRow] | list[dict[str, Any]], list[str]]:
    if period is None:
        return await _build_legacy_upload_rows(workbook, db, current_user)

    return await _build_period_upload_rows(workbook, db, current_user, period)


async def _load_upload_formula_context(db, users, period):
    formulas_by_user: dict[UUID, KPIFormula] = {}
    contracts_by_user: dict[UUID, frozenset[str]] = {}
    for user in users:
        try:
            formula = await _resolve_user_formula(db, user, period)
        except UploadValidationError:
            continue
        if formula is None:
            continue
        formulas_by_user[user.id] = formula
        try:
            contracts_by_user[user.id] = frozenset(
                field.key
                for field in build_header_contract([formula]).metric_headers
            )
        except HeaderContractError as exc:
            raise UploadValidationError(str(exc)) from exc
    try:
        contract = build_header_contract(list(formulas_by_user.values()))
    except HeaderContractError as exc:
        raise UploadValidationError(str(exc)) from exc
    return formulas_by_user, contracts_by_user, contract


async def _validate_period_upload_row(
    row_number,
    values,
    metadata_indexes,
    user_by_email,
    formulas_by_user,
    period,
    db,
    header_indexes,
    metric_by_key,
    contracts_by_user,
):
    email_value = values[metadata_indexes["email"]]
    email = (
        str(email_value).strip().casefold()
        if isinstance(email_value, str)
        else ""
    )
    if not email:
        return None, f"Row {row_number} — Email: a non-empty email is required."
    user = user_by_email.get(email)
    if user is None:
        return None, (
            f"Row {row_number} — Email '{email}': active employee was not found "
            "in your authorized departments."
        )
    formula = formulas_by_user.get(user.id)
    if formula is None:
        return None, (
            f"Row {row_number} — Email '{email}': no active KPI formula is assigned "
            "for the selected period."
        )
    existing_record = (
        await db.execute(
            select(KPIRecord).where(
                KPIRecord.user_id == user.id,
                KPIRecord.kpi_period_id == period.id,
                KPIRecord.is_deleted.is_(False),
            )
        )
    ).scalars().first()
    if existing_record and (existing_record.is_frozen or existing_record.is_overridden):
        return None, (
            f"Row {row_number} — Email '{email}': the existing KPI record is frozen "
            "or overridden and cannot be recalculated."
        )
    task_data, metric_errors = _read_upload_metrics(
        row_number,
        email,
        values,
        header_indexes,
        metric_by_key,
        contracts_by_user[user.id],
    )
    if metric_errors:
        return None, metric_errors
    return ValidatedUploadRow(
        row_number=row_number,
        email=email,
        user=user,
        formula=formula,
        task_data=task_data,
    ), None


async def _build_period_upload_rows(
    workbook: ParsedWorkbook,
    db: AsyncSession,
    current_user: User,
    period: KPIPeriod,
) -> tuple[list[ValidatedUploadRow], list[str]]:
    errors: list[str] = []
    department_ids = await get_authorized_department_ids(db, current_user)
    users_statement = _scoped_active_users_statement(department_ids)
    users = (await db.execute(users_statement)).scalars().unique().all()
    user_by_email = {
        str(user.email).strip().casefold(): user
        for user in users
        if user.email
    }
    formulas_by_user, contracts_by_user, contract = await _load_upload_formula_context(
        db, users, period
    )

    header_indexes, header_errors = _validate_headers(workbook.headers, contract)
    errors.extend(header_errors)
    if not workbook.rows:
        errors.append("The workbook contains no data rows.")
    if header_errors:
        return [], errors

    metadata_indexes = {
        key: header_indexes[normalize_header(header)]
        for header in EMPLOYEE_HEADERS
        for key in (normalize_header(header),)
    }
    metric_by_key = {
        field.key: field for field in contract.metric_headers
    }
    seen_emails: set[str] = set()
    validated_rows: list[ValidatedUploadRow] = []

    for row_number, values in workbook.rows:
        email_value = values[metadata_indexes["email"]]
        email = (
            str(email_value).strip().casefold()
            if isinstance(email_value, str)
            else ""
        )
        if not email:
            errors.append(f"Row {row_number} — Email: a non-empty email is required.")
            continue
        if email in seen_emails:
            errors.append(f"Row {row_number} — Email '{email}': duplicate employee row.")
            continue
        seen_emails.add(email)
        row, row_errors = await _validate_period_upload_row(
            row_number,
            values,
            metadata_indexes,
            user_by_email,
            formulas_by_user,
            period,
            db,
            header_indexes,
            metric_by_key,
            contracts_by_user,
        )
        if isinstance(row_errors, list):
            errors.extend(row_errors)
        elif row_errors:
            errors.append(row_errors)
        if row:
            validated_rows.append(row)
    return validated_rows, errors


async def _build_legacy_upload_rows(
    dataframe: Any, db: AsyncSession, current_user: User
) -> tuple[list[dict[str, Any]], list[str]]:
    department_ids = await get_authorized_department_ids(db, current_user)
    users = (
        await db.execute(_scoped_active_users_statement(department_ids))
    ).scalars().all()
    user_by_email = {
        str(email).strip().casefold(): user_id
        for user_id, email in users
        if email
    }
    rows: list[dict[str, Any]] = []
    errors: list[str] = []
    for row_number, (_, values) in enumerate(dataframe.iterrows(), start=2):
        email_value = values.get("Email")
        email = str(email_value).strip().casefold() if email_value else ""
        user_id = user_by_email.get(email)
        if user_id is None:
            errors.append(
                f"Row {row_number}: Active user with email '{email}' not found."
            )
            continue
        rows.append({"user_id": user_id})
    return rows, errors


def _scoped_active_users_statement(
    department_ids: set[str] | User | None,
    authorized_department_ids: set[str] | None = None,
):
    if hasattr(department_ids, "role"):
        department_ids = authorized_department_ids
    statement = select(User).options(
        selectinload(User.user_departments).selectinload(UserDepartment.department),
        selectinload(User.position),
    ).where(User.is_active.is_(True), User.is_deleted.is_(False))
    if department_ids is not None:
        database_department_ids = [UUID(department_id) for department_id in department_ids]
        statement = statement.where(
            User.id.in_(
                select(UserDepartment.user_id).where(
                    UserDepartment.department_id.in_(database_department_ids),
                    UserDepartment.is_deleted.is_(False),
                )
            )
        )
    return statement.order_by(func.lower(User.email))


def _visible_department_name(
    user: User, department_ids: set[str] | None
) -> str | None:
    for membership in user.user_departments or []:
        if membership.is_deleted:
            continue
        if (
            department_ids is not None
            and str(membership.department_id) not in department_ids
        ):
            continue
        if membership.department is not None:
            return membership.department.name
    return None


def _user_metadata(user: User, department_ids: set[str] | None) -> dict[str, str]:
    memberships = [
        membership
        for membership in (user.user_departments or [])
        if not membership.is_deleted
        and (
            department_ids is None
            or str(membership.department_id) in department_ids
        )
    ]
    return {
        "Email": str(user.email),
        "Employee Name": f"{user.first_name} {user.last_name}".strip(),
        "Department": _visible_department_name(user, department_ids) or "N/A",
        "Position": user.position.name if user.position else "N/A",
        "Level": memberships[0].user_level if memberships else "N/A",
    }


async def list_import_periods(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(admin_hr_checker),
):
    enforce_permission(current_user, "kpi.create")
    periods = (
        await db.execute(
            select(KPIPeriod)
            .where(
                KPIPeriod.is_deleted.is_(False),
                KPIPeriod.status.in_(EDITABLE_PERIOD_STATUSES),
            )
            .order_by(KPIPeriod.start_date.desc(), KPIPeriod.id)
        )
    ).scalars().all()
    return [
        {
            "id": str(period.id),
            "label": period.label,
            "month": _period_month_range(period).start_inclusive.strftime("%Y-%m"),
        }
        for period in periods
        if _period_month_range(period) is not None
    ]


@router.get("/periods", summary="List evaluation periods available for KPI import")
@limiter.limit("30/minute")
async def get_import_periods(
    request: Request,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(admin_hr_checker),
):
    return await list_import_periods(db, current_user)


@router.get("/export-template", summary="Download the official KPI import template")
@limiter.limit("10/minute")
async def export_kpi_template(
    request: Request,
    period_id: UUID = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(admin_hr_checker),
):
    enforce_permission(current_user, "report.export")
    try:
        period = _assert_importable_period(await db.get(KPIPeriod, period_id))
        department_ids = await get_authorized_department_ids(db, current_user)
        users = (
            await db.execute(_scoped_active_users_statement(department_ids))
        ).scalars().unique().all()
        formulas_by_user = {}
        for user in users:
            formula = await _resolve_user_formula(db, user, period)
            if formula is not None:
                formulas_by_user[user.id] = formula
        if not formulas_by_user:
            raise UploadValidationError(
                "No active employees have an applicable KPI formula for this period."
            )
        try:
            contract = build_header_contract(list(formulas_by_user.values()))
        except HeaderContractError as exc:
            raise UploadValidationError(str(exc)) from exc

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "KPI_Data"
        worksheet.append(contract.headers)
        for user in users:
            if user.id not in formulas_by_user:
                continue
            metadata = _user_metadata(user, department_ids)
            worksheet.append(
                [metadata.get(header, None) if header in EMPLOYEE_HEADERS else None
                 for header in contract.headers]
            )
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column_index, header in enumerate(contract.headers, start=1):
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(len(header) + 4, 18), 42
            )
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        output.seek(0)
        return StreamingResponse(
            output,
            headers={
                "Content-Disposition": 'attachment; filename="KPI_Master_Template.xlsx"'
            },
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    except UploadValidationError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail={"code": "validation_error", "errors": exc.errors},
        ) from exc
    except Exception:
        logger.exception("Failed to generate KPI upload template.")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail={
                "code": "processing_error",
                "message": "The server could not generate the template.",
            },
        ) from None


@router.post("/monthly-data", status_code=status.HTTP_200_OK)
@limiter.limit("3/minute")
async def upload_monthly_kpi_data(
    request: Request,
    file: UploadFile = File(...),
    period_id: Annotated[UUID | None, Form()] = None,
    target_month: Annotated[str | None, Form()] = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(admin_hr_checker),
):
    enforce_permission(current_user, "kpi.create")
    try:
        workbook = await _read_upload_dataframe(file)
        period = await _resolve_period(
            db, period_id=period_id, target_month=target_month
        )
        valid_rows, errors = await _build_upload_rows(
            workbook, db, current_user, period
        )
        if errors:
            raise UploadValidationError(errors)
        if not valid_rows:
            raise UploadValidationError("The workbook contains no valid employee rows.")

        period_range = _period_month_range(period)
        for row in valid_rows:
            await KPICalculationEngine.execute_kpi_calculation(
                db=db,
                user=row.user,
                period_start=period_range.start_inclusive,
                period_end=period_range.end_exclusive,
                period_type=PeriodType.MONTHLY,
                task_data=row.task_data,
                commit=False,
            )
            await db.flush()
            result = (
                await db.execute(
                    select(KPIRecord).where(
                        KPIRecord.user_id == row.user.id,
                        KPIRecord.kpi_period_id == period.id,
                        KPIRecord.period_type == PeriodType.MONTHLY,
                        KPIRecord.is_deleted.is_(False),
                    )
                )
            ).scalars().first()
            if result is None or result.formula_id != row.formula.id:
                raise RuntimeError(
                    f"KPI calculation did not persist expected context for {row.email}."
                )
        await db.commit()
        return {
            "message": "KPI data was imported and calculated successfully.",
            "processed_count": len(valid_rows),
            "status": "completed",
            "period_id": str(period.id),
        }
    except UploadValidationError as exc:
        await db.rollback()
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail={"code": "validation_error", "errors": exc.errors},
        ) from exc
    except Exception:
        await db.rollback()
        logger.exception("KPI workbook import failed unexpectedly.")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail={
                "code": "processing_error",
                "message": "The server could not process this file. Please try again.",
            },
        ) from None
