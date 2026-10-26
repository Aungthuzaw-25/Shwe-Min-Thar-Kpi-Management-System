import asyncio
import io
from datetime import UTC, date, datetime
from decimal import Decimal
from types import SimpleNamespace
from uuid import uuid4
from zipfile import ZipFile

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.future import select

from app.core.dependencies import get_current_user
from app.database.connection import Base, get_db
from app.database.seed import _ensure_active_period
from app.modules.departments.models import Department, UserDepartment
from app.modules.kpi.engine import KPICalculationEngine
from app.modules.kpi.models import (
    FormulaScope,
    FormulaStatus,
    KPIAssignment,
    KPIFormula,
    KPIFormulaPeriod,
    KPIHistory,
    KPIPeriod,
    KPIPeriodStatus,
    KPIRecord,
    KPIScoringRuleItem,
    PeriodType,
)
from app.modules.kpi.period_ranges import calendar_period_range
from app.modules.kpi.upload_contract import HeaderContractError, build_header_contract
from app.modules.kpi.upload_router import limiter, router
from app.modules.tasks.models.task import Task
from app.modules.tasks.models.task_review import TaskReview
from app.modules.users.models import Position, Role, User

TABLES = [
    Department.__table__,
    UserDepartment.__table__,
    Role.__table__,
    Position.__table__,
    User.__table__,
    KPIFormula.__table__,
        KPIFormulaPeriod.__table__,
    KPIPeriod.__table__,
    KPIAssignment.__table__,
    KPIRecord.__table__,
    KPIHistory.__table__,
    KPIScoringRuleItem.__table__,
    Task.__table__,
    TaskReview.__table__,
]


async def _setup(tmp_path):
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'kpi-upload.db'}")
    async with engine.begin() as connection:
        await connection.run_sync(
            lambda sync_connection: Base.metadata.create_all(
                sync_connection, tables=TABLES
            )
        )
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with session_factory() as db:
        department = Department(name=f"Upload department {uuid4()}")
        admin_role = Role(name="Super Admin")
        employee_role = Role(name=f"Employee {uuid4()}")
        db.add_all([department, admin_role, employee_role])
        await db.flush()

        position = Position(name=f"Upload position {uuid4()}", department_id=department.id)
        admin = User(
            email=f"admin-{uuid4()}@example.test",
            hashed_password=uuid4().hex,
            first_name="QA",
            last_name="Admin",
            role_id=admin_role.id,
            role=admin_role,
            is_active=True,
            is_deleted=False,
        )
        employees = [
            User(
                email=f"employee-{index}-{uuid4()}@example.test",
                hashed_password=uuid4().hex,
                first_name="QA",
                last_name=f"Employee {index}",
                role_id=employee_role.id,
                position=position,
                is_active=True,
                is_deleted=False,
                user_departments=[
                    UserDepartment(
                        department_id=department.id,
                        user_level="L2",
                        is_deleted=False,
                    )
                ],
            )
            for index in (1, 2)
        ]
        db.add_all([position, admin, *employees])
        await db.flush()

        current_month = date.today().replace(day=1)
        period_range = calendar_period_range(PeriodType.MONTHLY, current_month)
        period = KPIPeriod(
            label=f"Upload test {current_month:%Y-%m}",
            start_date=period_range.start_inclusive.astimezone(UTC).replace(tzinfo=None),
            end_date=period_range.end_exclusive.astimezone(UTC).replace(tzinfo=None),
            status=KPIPeriodStatus.OPEN,
            is_deleted=False,
        )
        formula = KPIFormula(
            name="Upload regression formula",
            department_id=department.id,
            position_name=position.name,
            position_id=position.id,
            user_level="L2",
            categories=[
                {
                    "name": "Attendance",
                    "weight": 1.0,
                    "max_score": 5.0,
                    "evaluation_type": "DEDUCTION",
                    "deductions": [
                        {
                            "name": "Late Arrival",
                            "measure_type": "count",
                            "rules": [{"operator": ">=", "value": 2, "deduct_score": 2}],
                        }
                    ],
                }
            ],
            scope=FormulaScope.DEPARTMENT,
            is_active=True,
            status=FormulaStatus.ACTIVE,
            effective_month=current_month.strftime("%Y-%m"),
            version=1,
            is_deleted=False,
            scoring_rules=[
                KPIScoringRuleItem(
                    category_name="attendance",
                    rule_name="late arrival",
                    measure_type="count",
                    operator=">=",
                    target_value=2,
                    deduct_score=2,
                )
            ],
        )
        db.add_all([period, formula])
        await db.commit()
        return engine, session_factory, admin, employees, period, formula


async def _client(session_factory, admin):
    from fastapi import FastAPI

    app = FastAPI()
    app.include_router(router, prefix="/api/v1")
    app.state.limiter = limiter
    app.state.limiter.enabled = False

    async def override_db():
        async with session_factory() as db:
            yield db

    app.dependency_overrides[get_db] = override_db
    # E2: upload routes are gated by get_current_user + permission
    # keys now (the old admin_hr_checker role gate is deleted).
    app.dependency_overrides[get_current_user] = lambda: admin
    return AsyncClient(transport=ASGITransport(app=app), base_url="http://test")


async def _template(client, period_id):
    response = await client.get(
        "/api/v1/kpi-uploads/export-template",
        params={"period_id": str(period_id)},
    )
    assert response.status_code == 200, response.text
    assert 'filename="KPI_Master_Template.xlsx"' in response.headers["content-disposition"]
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    return response.content


def _workbook_bytes(content, edit):
    workbook = load_workbook(io.BytesIO(content))
    worksheet = workbook["KPI_Data"]
    edit(workbook, worksheet)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


async def _upload(client, content, period_id, filename="KPI_Master_Template.xlsx"):
    return await client.post(
        "/api/v1/kpi-uploads/monthly-data",
        data={"period_id": str(period_id)} if period_id else {},
        files={
            "file": (
                filename,
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )


def test_official_template_round_trip_calculates_period_records(tmp_path):
    async def scenario():
        engine, factory, admin, employees, period, formula = await _setup(tmp_path)
        try:
            async with await _client(factory, admin) as client:
                periods = await client.get("/api/v1/kpi-uploads/periods")
                assert periods.status_code == 200
                assert any(item["id"] == str(period.id) for item in periods.json())
                content = await _template(client, period.id)
                workbook = load_workbook(io.BytesIO(content), data_only=False)
                assert workbook.sheetnames == ["KPI_Data"]
                worksheet = workbook["KPI_Data"]
                headers = [cell.value for cell in worksheet[1]]
                assert headers == [
                    "Email",
                    "Employee Name",
                    "Department",
                    "Position",
                    "Level",
                    "Late Arrival",
                ]
                assert worksheet.max_row == len(employees) + 1
                assert not worksheet.merged_cells.ranges
                assert not worksheet.data_validations.dataValidation
                assert worksheet.sheet_state == "visible"
                assert all(
                    not dimension.hidden
                    for dimension in worksheet.column_dimensions.values()
                )
                worksheet.cell(1, 1, " email ")
                worksheet.cell(1, 6, " late   ARRIVAL ")
                for row_index, metric in enumerate((1.25, "2"), start=2):
                    worksheet.cell(row_index, 6, metric)
                round_trip_path = tmp_path / "KPI_Master_Template_generated.xlsx"
                workbook.save(round_trip_path)
                workbook.close()

                reopened = load_workbook(round_trip_path, data_only=False)
                assert [reopened["KPI_Data"].cell(row, 6).value for row in (2, 3)] == [
                    1.25,
                    "2",
                ]
                reopened.close()

                response = await _upload(
                    client, round_trip_path.read_bytes(), period.id
                )
                assert response.status_code == 200, response.text
                assert response.json()["processed_count"] == 2
                assert response.json()["period_id"] == str(period.id)
                repeated_response = await _upload(
                    client, round_trip_path.read_bytes(), period.id
                )
                assert repeated_response.status_code == 200, repeated_response.text

            async with factory() as db:
                records = (
                    await db.execute(select(KPIRecord).order_by(KPIRecord.user_id))
                ).scalars().all()
                assert len(records) == 2
                assert all(record.kpi_period_id == period.id for record in records)
                assert all(record.formula_id == formula.id for record in records)
                assert sorted(record.percentage for record in records) == [60.0, 100.0]
                history = (await db.execute(select(KPIHistory))).scalars().all()
                assert len(history) == 2
        finally:
            limiter.enabled = True
            await engine.dispose()

    asyncio.run(scenario())


def test_import_period_list_excludes_open_period_with_nonmonthly_boundaries(tmp_path):
    async def scenario():
        engine, factory, admin, _employees, valid_period, _formula = await _setup(tmp_path)
        invalid_period_id = uuid4()
        try:
            async with factory() as db:
                db.add(
                    KPIPeriod(
                        id=invalid_period_id,
                        label="testing",
                        start_date=datetime(2026, 10, 20, 17, 30),
                        end_date=datetime(2026, 10, 22, 17, 30),
                        status=KPIPeriodStatus.OPEN,
                        is_deleted=False,
                    )
                )
                await db.commit()

            async with await _client(factory, admin) as client:
                response = await client.get("/api/v1/kpi-uploads/periods")
                assert response.status_code == 200
                listed_ids = {item["id"] for item in response.json()}
                assert str(valid_period.id) in listed_ids
                assert str(invalid_period_id) not in listed_ids

                export = await client.get(
                    "/api/v1/kpi-uploads/export-template",
                    params={"period_id": str(invalid_period_id)},
                )
                assert export.status_code == 400
                assert "complete monthly boundaries" in str(export.json()["detail"])
        finally:
            limiter.enabled = True
            await engine.dispose()

    asyncio.run(scenario())


def test_template_endpoint_rejects_missing_malformed_and_unknown_period_ids(tmp_path):
    async def scenario():
        engine, factory, admin, _employees, _period, _formula = await _setup(tmp_path)
        try:
            async with await _client(factory, admin) as client:
                missing = await client.get("/api/v1/kpi-uploads/export-template")
                malformed = await client.get(
                    "/api/v1/kpi-uploads/export-template",
                    params={"period_id": "not-a-uuid"},
                )
                unknown = await client.get(
                    "/api/v1/kpi-uploads/export-template",
                    params={"period_id": str(uuid4())},
                )
                assert missing.status_code == 422
                assert malformed.status_code == 422
                assert unknown.status_code == 400
                assert "available evaluation period" in str(unknown.json()["detail"])
        finally:
            limiter.enabled = True
            await engine.dispose()

    asyncio.run(scenario())


def test_template_endpoint_reports_period_without_applicable_formula(tmp_path):
    async def scenario():
        engine, factory, admin, _employees, period, formula = await _setup(tmp_path)
        try:
            async with factory() as db:
                persisted_formula = await db.get(KPIFormula, formula.id)
                persisted_formula.is_active = False
                await db.commit()

            async with await _client(factory, admin) as client:
                response = await client.get(
                    "/api/v1/kpi-uploads/export-template",
                    params={"period_id": str(period.id)},
                )
                assert response.status_code == 400
                # Split error (2026-10-24): the formula side is the
                # problem here (the only formula is inactive), so the
                # message must say no formula applies — not blame
                # employees.
                assert "No KPI formula applies to this period yet" in str(
                    response.json()["detail"]
                )
        finally:
            limiter.enabled = True
            await engine.dispose()

    asyncio.run(scenario())


def test_import_period_endpoints_reject_unauthenticated_and_unauthorized_users(
    tmp_path,
):
    async def scenario():
        engine, factory, _admin, _employees, _period, _formula = await _setup(tmp_path)
        from fastapi import FastAPI

        app = FastAPI()
        app.include_router(router, prefix="/api/v1")
        app.state.limiter = limiter

        async def override_db():
            async with factory() as db:
                yield db

        app.dependency_overrides[get_db] = override_db
        limiter.enabled = False
        try:
            async with AsyncClient(
                transport=ASGITransport(app=app), base_url="http://test"
            ) as client:
                unauthenticated = await client.get(
                    "/api/v1/kpi-uploads/periods"
                )
                assert unauthenticated.status_code == 401

                app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(
                    role=SimpleNamespace(name="Employee")
                )
                unauthorized = await client.get("/api/v1/kpi-uploads/periods")
                assert unauthorized.status_code == 403
        finally:
            limiter.enabled = True
            await engine.dispose()

    asyncio.run(scenario())


def test_standard_seed_period_uses_canonical_month_boundaries(tmp_path):
    async def scenario():
        engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'seed-period.db'}")
        async with engine.begin() as connection:
            await connection.run_sync(
                lambda sync_connection: KPIPeriod.__table__.create(sync_connection)
            )
        factory = async_sessionmaker(engine, expire_on_commit=False)
        try:
            async with factory() as db:
                period = await _ensure_active_period(db)
                period_month = datetime.strptime(period.label, "%B %Y")
                expected = calendar_period_range(
                    PeriodType.MONTHLY,
                    date(period_month.year, period_month.month, 1),
                )
                assert period.start_date == expected.start_inclusive.replace(
                    tzinfo=None
                )
                assert period.end_date == expected.end_exclusive.replace(
                    tzinfo=None
                )
        finally:
            await engine.dispose()

    asyncio.run(scenario())


def test_header_builder_rejects_normalization_collisions():
    formula = SimpleNamespace(
        categories=[
            {
                "name": "Attendance",
                "evaluation_type": "DEDUCTION",
                "deductions": [
                    {"name": "Late Arrival", "rules": [{}]},
                    {"name": " late   arrival ", "rules": [{}]},
                ],
            }
        ],
        scoring_rules=[
            SimpleNamespace(rule_name="late arrival"),
        ],
    )
    with pytest.raises(HeaderContractError, match="normalization collision"):
        build_header_contract([formula])


@pytest.mark.parametrize(
    ("value", "valid"),
    [
        (1, True),
        (1.25, True),
        (Decimal("2.5"), True),
        ("3.5", True),
        (None, False),
        ("", False),
        ("  ", False),
        ("NaN", False),
        ("Infinity", False),
        ("abc", False),
    ],
)
def test_excel_metric_number_normalization_is_finite(value, valid):
    from app.modules.kpi.upload_router import _finite_number

    result = _finite_number(value)
    assert (result is not None) is valid
    if valid:
        assert isinstance(result, float)


@pytest.mark.parametrize(
    ("case", "status_code"),
    [
        ("missing_column", 400),
        ("unknown_column", 400),
        ("duplicate_header", 400),
        ("blank_email", 400),
        ("blank_metric", 400),
        ("whitespace_metric", 400),
        ("unknown_employee", 400),
        ("duplicate_employee", 400),
        ("invalid_number", 400),
        ("nan", 400),
        ("infinity", 400),
        ("formula", 400),
        ("empty_workbook", 400),
        ("header_only", 400),
        ("corrupted_xlsx", 400),
        ("unsupported_extension", 400),
        ("missing_period", 400),
        ("invalid_period", 400),
        ("rule_header_mismatch", 400),
        ("mixed_valid_invalid_rows", 400),
    ],
)
def test_invalid_workbooks_return_controlled_validation_and_write_nothing(
    tmp_path, case, status_code
):
    async def scenario():
        engine, factory, admin, employees, period, _formula = await _setup(tmp_path)
        try:
            async with await _client(factory, admin) as client:
                official = await _template(client, period.id)

                def edit(workbook, sheet):
                    if case == "missing_column":
                        sheet.delete_cols(6)
                    elif case == "unknown_column":
                        sheet.cell(1, 6, "Not A KPI Rule")
                    elif case == "duplicate_header":
                        sheet.cell(1, 6, " email ")
                    elif case == "blank_email":
                        sheet.cell(2, 1, " ")
                    elif case == "blank_metric":
                        sheet.cell(2, 6, None)
                    elif case == "whitespace_metric":
                        sheet.cell(2, 6, "   ")
                    elif case == "unknown_employee":
                        sheet.cell(2, 1, "unknown@example.test")
                    elif case == "duplicate_employee":
                        sheet.append([cell.value for cell in sheet[2]])
                    elif case == "invalid_number":
                        sheet.cell(2, 6, "abc")
                    elif case == "nan":
                        sheet.cell(2, 6, "NaN")
                    elif case == "infinity":
                        sheet.cell(2, 6, "Infinity")
                    elif case == "formula":
                        sheet.cell(2, 6, "=1+1")
                    elif case == "empty_workbook":
                        sheet.delete_rows(1, sheet.max_row)
                    elif case == "header_only":
                        sheet.delete_rows(2, sheet.max_row)
                    elif case == "rule_header_mismatch":
                        sheet.cell(1, 6, "Unrelated Rule")
                    elif case == "mixed_valid_invalid_rows":
                        sheet.cell(2, 6, 1)
                        sheet.cell(3, 6, "invalid")

                if case in {"corrupted_xlsx", "unsupported_extension"}:
                    if case == "corrupted_xlsx":
                        invalid_file = io.BytesIO()
                        with ZipFile(invalid_file, "w"):
                            pass
                        invalid = invalid_file.getvalue()
                    else:
                        invalid = b"not a workbook"
                else:
                    invalid = _workbook_bytes(official, edit)
                if case == "missing_period":
                    response = await _upload(client, invalid, None)
                elif case == "invalid_period":
                    response = await _upload(client, invalid, uuid4())
                elif case == "unsupported_extension":
                    response = await _upload(client, official, period.id, "workbook.xls")
                else:
                    response = await _upload(client, invalid, period.id)
                assert response.status_code == status_code, response.text
                detail = response.json()["detail"]
                assert detail["code"] == "validation_error"
                assert detail["errors"]
                expected_fragments = {
                    "missing_column": "Missing required column",
                    "unknown_column": "Unknown column",
                    "duplicate_header": "Duplicate column header",
                    "blank_email": "non-empty email is required",
                    "blank_metric": "must be a finite number",
                    "whitespace_metric": "must be a finite number",
                    "unknown_employee": "active employee was not found",
                    "duplicate_employee": "duplicate employee row",
                    "invalid_number": "must be a finite number",
                    "nan": "must be a finite number",
                    "infinity": "must be a finite number",
                    "formula": "formula cells are not supported",
                    "empty_workbook": "worksheet is empty",
                    "header_only": "no data rows",
                    "corrupted_xlsx": "corrupted or invalid",
                    "unsupported_extension": "Unsupported file format",
                    "missing_period": "Select an evaluation period",
                    "invalid_period": "Select an available evaluation period",
                    "rule_header_mismatch": "Missing required column",
                    "mixed_valid_invalid_rows": "must be a finite number",
                }
                assert any(
                    expected_fragments[case] in error for error in detail["errors"]
                ), detail["errors"]

            async with factory() as db:
                assert (await db.execute(select(KPIRecord))).scalars().all() == []
                assert (await db.execute(select(KPIHistory))).scalars().all() == []
        finally:
            limiter.enabled = True
            await engine.dispose()

    asyncio.run(scenario())


def test_calculation_failure_rolls_back_prior_rows(tmp_path, monkeypatch):
    async def scenario():
        engine, factory, admin, employees, period, _formula = await _setup(tmp_path)
        original = KPICalculationEngine.execute_kpi_calculation
        calls = 0

        async def fail_on_second_row(*args, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise RuntimeError("injected calculation failure")
            return await original(*args, **kwargs)

        monkeypatch.setattr(
            KPICalculationEngine, "execute_kpi_calculation", fail_on_second_row
        )
        try:
            async with await _client(factory, admin) as client:
                official = await _template(client, period.id)
                valid = _workbook_bytes(
                    official,
                    lambda _workbook, sheet: (
                        sheet.cell(2, 6, 1),
                        sheet.cell(3, 6, 2),
                    ),
                )
                response = await _upload(client, valid, period.id)
                assert response.status_code == 500
                assert response.json()["detail"]["code"] == "processing_error"
            async with factory() as db:
                assert (await db.execute(select(KPIRecord))).scalars().all() == []
                assert (await db.execute(select(KPIHistory))).scalars().all() == []
        finally:
            limiter.enabled = True
            await engine.dispose()

    asyncio.run(scenario())
